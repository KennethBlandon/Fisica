"""
importadores
────────────
Lectura de datos para el módulo Circuito desde texto, CSV o Excel.
"""

from __future__ import annotations

from pathlib import Path
import csv

from .calculos import parse_prefixed_value


def _limpiar_tokens(texto: str):
    for linea in texto.splitlines():
        linea = linea.strip()
        if not linea or linea.startswith("#"):
            continue
        linea = linea.replace(";", ",").replace("\t", ",")
        for parte in linea.split(","):
            parte = parte.strip()
            if not parte:
                continue
            if "=" in parte:
                parte = parte.split("=", 1)[1].strip()
            yield parte


def cargar_desde_texto(ruta: str | Path):
    ruta = Path(ruta)
    tokens = list(_limpiar_tokens(ruta.read_text(encoding="utf-8")))
    if len(tokens) < 2:
        raise ValueError("El archivo debe tener una fuente y al menos un valor.")
    fuente = parse_prefixed_value(tokens[0])
    valores = [parse_prefixed_value(token) for token in tokens[1:]]
    return fuente, valores


def cargar_desde_csv(ruta: str | Path):
    ruta = Path(ruta)
    tokens = []
    with ruta.open(newline="", encoding="utf-8-sig") as archivo:
        lector = csv.reader(archivo)
        for fila in lector:
            for celda in fila:
                celda = celda.strip()
                if not celda or celda.startswith("#"):
                    continue
                if "=" in celda:
                    celda = celda.split("=", 1)[1].strip()
                tokens.append(celda)
    if len(tokens) < 2:
        raise ValueError("El CSV debe tener una fuente y al menos un valor.")
    fuente = parse_prefixed_value(tokens[0])
    valores = [parse_prefixed_value(token) for token in tokens[1:]]
    return fuente, valores


def cargar_desde_excel(ruta: str | Path):
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise ImportError(
            "Para leer Excel instala openpyxl: pip install openpyxl"
        ) from exc

    ruta = Path(ruta)
    wb = load_workbook(ruta, data_only=True)
    ws = wb.active

    celdas = []
    for fila in ws.iter_rows(values_only=True):
        for celda in fila:
            if celda is None:
                continue
            texto = str(celda).strip()
            if not texto or texto.startswith("#"):
                continue
            if "=" in texto:
                texto = texto.split("=", 1)[1].strip()
            celdas.append(texto)

    if len(celdas) < 2:
        raise ValueError("El Excel debe tener una fuente y al menos un valor.")

    fuente = parse_prefixed_value(celdas[0])
    valores = [parse_prefixed_value(token) for token in celdas[1:]]
    return fuente, valores


def cargar_desde_archivo(ruta: str | Path):
    ruta = Path(ruta)
    extension = ruta.suffix.lower()
    if extension in {".txt", ".dat"}:
        return cargar_desde_texto(ruta)
    if extension == ".csv":
        return cargar_desde_csv(ruta)
    if extension in {".xlsx", ".xlsm"}:
        return cargar_desde_excel(ruta)
    raise ValueError("Formato no soportado. Usa .txt, .csv, .xlsx o .xlsm.")
